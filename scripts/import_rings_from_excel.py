import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

# Ensure parent directory is in path to import database module
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import database
import openpyxl

excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'FOC Route IP All Ring UID BANTEN.xlsx')
print(f"Loading Excel file from: {excel_path} ...")

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['FINAL SUBMITTED']

client = database.supabase_client
mode = database.get_storage_mode()
print(f"Target Database Storage Mode: {mode}")

if mode != 'SUPABASE' or not client:
    print("Error: Supabase is not configured or not active.")
    sys.exit(1)

# Fetch existing ring codes in Supabase
existing_rings = client.table('rings').select('id, code, name').execute().data or []
existing_codes = {r['code'].strip().upper(): r for r in existing_rings}
print(f"Existing rings in database ({len(existing_rings)}): {list(existing_codes.keys())}")

imported_count = 0
skipped_count = 0

print("\n" + "=" * 65)
print("STARTING IMPORT PROCESS (RINGS 6 TO 22)")
print("=" * 65)

# Rows in Excel:
# Row 2 = Ring 1
# Row 3 = Ring 2
# ...
# Row 7 = Ring 6
# Row 23 = Ring 22
for r in range(7, 24):
    row_no = ws.cell(r, 1).value
    if row_no is None:
        continue

    ring_code = f"RING {row_no}"
    uid = str(ws.cell(r, 2).value or 'BANTEN').strip()
    foc_line = str(ws.cell(r, 5).value or f"AMI_RJKT_UP3CIKOKOL_RING_{row_no}").strip()
    pop1_loc = str(ws.cell(r, 6).value or '').strip()
    pop2_loc = str(ws.cell(r, 48).value or '').strip()
    total_len_excel = ws.cell(r, 51).value

    # Skip if already exists
    if ring_code.upper() in existing_codes:
        print(f"[SKIP] {ring_code} already exists in database (ID: {existing_codes[ring_code.upper()]['id']}). Keeping untouched.")
        skipped_count += 1
        continue

    print(f"\n--- Importing {ring_code} ({foc_line}) ---")
    print(f"  POP 1: {pop1_loc}")
    print(f"  POP 2: {pop2_loc}")

    # 1. Insert Ring
    ring_res = client.table('rings').insert({
        'name': uid,
        'code': ring_code,
        'description': foc_line
    }).execute()
    new_ring_id = ring_res.data[0]['id']
    print(f"  -> Created Ring ID: {new_ring_id}")

    # 2. Collect DCU codes and lengths
    # DCU k is at col 10 + (k-1)*4
    dcu_codes = []
    for k in range(1, 11):
        dcu_col = 10 + (k - 1) * 4
        dcu_val = ws.cell(r, dcu_col).value
        if dcu_val is not None and str(dcu_val).strip() != '':
            dcu_codes.append(str(dcu_val).strip())
        else:
            break

    # 3. Insert Nodes (POP 1, DCUs, POP 2)
    nodes_payload = []
    # POP 1
    nodes_payload.append({
        'ring_id': new_ring_id,
        'node_type': 'POP_START',
        'name': 'POP 1 (MULAI)',
        'code': f"GARDU POP-START-{new_ring_id}",
        'location': pop1_loc,
        'coordinates': None,
        'detail_json': {},
        'position_order': 1
    })

    # DCUs
    for idx, dcu_code in enumerate(dcu_codes, start=1):
        clean_dcu_code = dcu_code if dcu_code.upper().startswith('GARDU') else f"GARDU {dcu_code}"
        nodes_payload.append({
            'ring_id': new_ring_id,
            'node_type': 'DCU',
            'name': f"DCU {idx}",
            'code': clean_dcu_code,
            'location': 'Jalur Lintasan DCU',
            'coordinates': None,
            'detail_json': {},
            'position_order': idx + 1
        })

    # POP 2
    nodes_payload.append({
        'ring_id': new_ring_id,
        'node_type': 'POP_END',
        'name': 'POP 2 (AKHIR)',
        'code': f"GARDU POP-END-{new_ring_id}",
        'location': pop2_loc,
        'coordinates': None,
        'detail_json': {},
        'position_order': len(dcu_codes) + 2
    })

    nodes_res = client.table('nodes').insert(nodes_payload).execute()
    created_nodes = sorted(nodes_res.data, key=lambda n: n['position_order'])
    print(f"  -> Created {len(created_nodes)} Nodes (POP 1, {len(dcu_codes)} DCUs, POP 2)")

    # 4. Insert Segments
    # Total segments = len(dcu_codes) + 1
    segments_payload = []
    
    # Segment 1: POP 1 -> DCU 1 (Col 7)
    l1 = ws.cell(r, 7).value
    try:
        dist_1 = float(l1) if l1 is not None else 2000.0
    except Exception:
        dist_1 = 2000.0

    segments_payload.append({
        'ring_id': new_ring_id,
        'from_node_id': created_nodes[0]['id'],
        'to_node_id': created_nodes[1]['id'],
        'distance_m': dist_1,
        'status': 'AMAN',
        'cable_type': 'ADSS 24 Core',
        'notes': 'Segmen awal',
        'position_order': 1
    })

    # Middle segments: DCU i -> DCU i+1
    for i in range(1, len(dcu_codes)):
        len_col = 7 + i * 4
        dist_val = ws.cell(r, len_col).value
        try:
            d_float = float(dist_val) if dist_val is not None else 2000.0
        except Exception:
            d_float = 2000.0

        segments_payload.append({
            'ring_id': new_ring_id,
            'from_node_id': created_nodes[i]['id'],
            'to_node_id': created_nodes[i + 1]['id'],
            'distance_m': d_float,
            'status': 'AMAN',
            'cable_type': 'ADSS 24 Core',
            'notes': '',
            'position_order': i + 1
        })

    # Last segment: DCU last -> POP 2
    last_dcu_idx = len(dcu_codes)
    last_len_col = 7 + last_dcu_idx * 4
    last_dist_val = ws.cell(r, last_len_col).value
    try:
        last_d_float = float(last_dist_val) if last_dist_val is not None else 2000.0
    except Exception:
        last_d_float = 2000.0

    segments_payload.append({
        'ring_id': new_ring_id,
        'from_node_id': created_nodes[last_dcu_idx]['id'],
        'to_node_id': created_nodes[last_dcu_idx + 1]['id'],
        'distance_m': last_d_float,
        'status': 'AMAN',
        'cable_type': 'ADSS 24 Core',
        'notes': '',
        'position_order': last_dcu_idx + 1
    })

    segs_res = client.table('segments').insert(segments_payload).execute()
    total_dist = sum(s['distance_m'] for s in segs_res.data)
    print(f"  -> Created {len(segs_res.data)} Segments. Total Distance: {total_dist:,.0f} m (Excel Total: {total_len_excel} m)")
    imported_count += 1

print("\n" + "=" * 65)
print(f"IMPORT COMPLETE: {imported_count} rings imported, {skipped_count} rings skipped.")
print("=" * 65)
